from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from docx import Document

from test_case_agent.source_json_projection import build_docx_source_json


class SourceJsonProjectionTests(unittest.TestCase):
    def test_docx_projection_preserves_order_sections_and_table_rows(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            docx = root / "source" / "main.docx"
            docx.parent.mkdir(parents=True)
            document = Document()
            document.add_heading("4.3 Карточка заявки", level=1)
            document.add_paragraph("Перед таблицей.")
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Поле"
            table.cell(0, 1).text = "Ограничение"
            table.cell(1, 0).text = "Фамилия"
            table.cell(1, 1).text = "Текстовые символы и дефис"
            document.save(docx)

            first = build_docx_source_json(docx, repo_root=root)
            second = build_docx_source_json(docx, repo_root=root)

        self.assertEqual(first["projection_sha256"], second["projection_sha256"])
        self.assertEqual("source/main.docx", first["source_path"])
        self.assertEqual(4, first["block_count"])
        self.assertEqual(
            ["heading-1", "paragraph", "table-row", "table-row"],
            [item["kind"] for item in first["blocks"]],
        )
        self.assertEqual(
            ["Фамилия", "Текстовые символы и дефис"],
            first["blocks"][3]["cells"],
        )
        self.assertEqual(
            ["4.3 Карточка заявки"],
            first["blocks"][3]["section_path"],
        )


if __name__ == "__main__":
    unittest.main()
