from __future__ import annotations

import hashlib
import json
import shutil
import unittest
from dataclasses import replace
from pathlib import Path
from unittest.mock import patch

from docx import Document
from PIL import Image

from test_case_agent.coverage_graph import (
    CoverageGraph,
    CoverageObligation,
    CoverageProperty,
    build_coverage_graph,
)
from test_case_agent.coverage_io import load_property_derivations
from test_case_agent.iteration_contract import (
    reviewer_acceptance_contract,
    validate_suite,
)
from test_case_agent.review_cycle.prepared_package import (
    PreparedDictionaryFixtureProvenance,
    PreparedDictionaryRequirement,
    PreparedDictionaryValue,
    PreparedObligation,
    PreparedObligationSet,
    load_obligations,
    prepared_dictionary_value_set_sha256,
)
from test_case_agent.review_cycle.source_assertions import (
    RegisteredEvidenceSource,
    SupportingSourceBinding,
    build_source_assertion_manifest,
    parse_embedded_source_assertion_contract,
)
from test_case_agent.reviewer_evidence import (
    ReviewerEvidenceError,
    _qualified_closed_fixture_subset_provenance,
    build_design_support_mapping,
    build_reviewer_evidence_pack,
    _dictionary_slice,
    load_reviewer_evidence_basis_document,
    prepare_reviewer_evidence_basis,
)
from test_case_agent.scope_compiler import compile_scope_source
from test_case_agent.scope_registry import load_and_resolve_scope_registry
from test_case_agent.test_design import build_test_design_plan, render_test_cases


class ReviewerEvidenceTests(unittest.TestCase):
    def setUp(self) -> None:
        from tests.test_source_qualified_run import SourceQualifiedRunTests

        fixture = SourceQualifiedRunTests(
            methodName="test_offline_run_is_source_qualified_and_preserves_canonical"
        )
        fixture.setUp()
        self.addCleanup(fixture.doCleanups)
        self.fixture = fixture
        self.registry = load_and_resolve_scope_registry(
            fixture.registry,
            fixture.ft,
            scope_id="sample-scope",
        )
        self.compiled = compile_scope_source(
            self.registry,
            scope_id="sample-scope",
            repo_root=fixture.repo,
        )
        expected_ids = tuple(
            item.obligation_id
            for item in fixture.prepared_obligations.obligations
            if item.coverage_status == "testable"
        )
        self.contract = parse_embedded_source_assertion_contract(
            fixture.source_evidence.read_text(encoding="utf-8"),
            fixture.repo,
            expected_scope_slug="sample-scope",
            expected_obligation_ids=expected_ids,
        )
        derivations = load_property_derivations(
            fixture.derivations,
            expected_source_manifest_digest=self.contract.manifest.digest,
            expected_obligation_set_digest=fixture.prepared_obligations.digest,
        )
        self.graph = build_coverage_graph(
            ft_slug="sample",
            tc_prefix="SMP",
            source_manifest=self.contract.manifest,
            obligation_set=fixture.prepared_obligations,
            derivations=derivations.derivations,
        )
        self.basis = prepare_reviewer_evidence_basis(
            fixture.repo,
            self.compiled,
            self.contract.manifest,
            self.contract.review_receipt,
            fixture.prepared_obligations,
        )

    def _fixture_provenance(
        self,
        dictionary_id: str,
        values: tuple[PreparedDictionaryValue, ...],
        *,
        external_dynamic: bool = False,
    ) -> tuple[
        PreparedDictionaryFixtureProvenance,
        tuple[RegisteredEvidenceSource, ...],
    ]:
        root = self.fixture.ft / "work" / "dictionary-provenance"
        root.mkdir(parents=True, exist_ok=True)
        slug = dictionary_id.casefold().replace("-", "_")
        fixture_id = "FX-" + dictionary_id.removeprefix("DICT-")
        if not external_dynamic:
            source = root / f"{slug}.fixture.json"
            source.write_text(
                json.dumps(
                    {
                        "dictionary_id": dictionary_id,
                        "fixture_id": fixture_id,
                        "fixture_value_set_sha256": (
                            prepared_dictionary_value_set_sha256(values)
                        ),
                        "values": [item.value for item in values],
                    },
                    ensure_ascii=False,
                    sort_keys=True,
                    separators=(",", ":"),
                ),
                encoding="utf-8",
            )
            source_relative = self.fixture._relative(source)
            source_sha256 = hashlib.sha256(source.read_bytes()).hexdigest()
            return (
                PreparedDictionaryFixtureProvenance(
                    evidence_kind="registered-fixture-source",
                    dictionary_id=dictionary_id,
                    fixture_id=fixture_id,
                    source_path=source_relative,
                    source_locator="json-pointer:/values",
                    source_sha256=source_sha256,
                    value_set_sha256=prepared_dictionary_value_set_sha256(values),
                    version="not-declared",
                    effective_date="not-declared",
                ),
                (
                    RegisteredEvidenceSource(
                        source_relative,
                        source_sha256,
                        "supporting-material",
                    ),
                ),
            )

        source = root / f"{slug}.response.json"
        source.write_text(
            json.dumps(
                {"suggestions": [{"value": values[0].value}]},
                ensure_ascii=False,
                sort_keys=True,
                separators=(",", ":"),
            ),
            encoding="utf-8",
        )
        source_relative = self.fixture._relative(source)
        source_sha256 = hashlib.sha256(source.read_bytes()).hexdigest()
        verified_at = "2026-07-23T08:15:30Z"
        query = f"query-{fixture_id}"
        request = {
            "endpoint": "https://suggestions.dadata.ru/suggestions/api/4_1/rs/suggest/address",
            "method": "POST",
            "parameters": {"query": query},
        }
        expected_response = {
            "outcome": "suggestions-found",
            "exact_suggestion": values[0].value,
            "exact_components": {},
        }
        receipt = root / f"{slug}.verification.json"
        receipt.write_text(
            json.dumps(
                {
                    "schema_version": 1,
                    "evidence_kind": "verified-live-positive-response",
                    "fixture_id": fixture_id,
                    "status": "verified",
                    "provider": "DaData",
                    "request": request,
                    "response_snapshot": source.name,
                    "response_sha256": source_sha256,
                    "expected_response": expected_response,
                    "verification": {
                        "all_exact_suggestion_matched": True,
                        "all_http_200": True,
                        "attempt_count": 1,
                        "attempts": [
                            {
                                "attempt": 1,
                                "checked_at_utc": verified_at,
                                "http_status": 200,
                                "response_sha256": source_sha256,
                                "exact_suggestion_matched": True,
                            }
                        ],
                    },
                },
                ensure_ascii=False,
                sort_keys=True,
                separators=(",", ":"),
            ),
            encoding="utf-8",
        )
        receipt_relative = self.fixture._relative(receipt)
        receipt_sha256 = hashlib.sha256(receipt.read_bytes()).hexdigest()
        catalog = root / f"{slug}.fixture-catalog.md"
        catalog.write_text(
            "# Fixture Catalog\n\n"
            "| fixture_id | query | suggestion | sources | checked_at | response_sha256 | lifecycle | status |\n"
            "| --- | --- | --- | --- | --- | --- | --- | --- |\n"
            f"| `{fixture_id}` | `{query}` | `{values[0].value}` | "
            f"`{source.name}`; `{receipt.name}` | `{verified_at}` | "
            f"`{source_sha256}` | `verified-once / revalidate-on-failure` | "
            "`verified-live-response` |\n",
            encoding="utf-8",
        )
        catalog_relative = self.fixture._relative(catalog)
        catalog_sha256 = hashlib.sha256(catalog.read_bytes()).hexdigest()
        return (
            PreparedDictionaryFixtureProvenance(
                evidence_kind="external-dynamic-fixture",
                dictionary_id=dictionary_id,
                fixture_id=fixture_id,
                source_path=source_relative,
                source_locator="json-pointer:",
                source_sha256=source_sha256,
                value_set_sha256=prepared_dictionary_value_set_sha256(values),
                version="1",
                effective_date="2026-07-23",
                provider="DaData",
                verification_receipt_path=receipt_relative,
                verification_receipt_locator="json-pointer:",
                verification_receipt_sha256=receipt_sha256,
                lifecycle_source_path=catalog_relative,
                lifecycle_source_locator=f"markdown-table-row:{fixture_id}",
                lifecycle_source_sha256=catalog_sha256,
                verified_at=verified_at,
                lifecycle_policy="verified-once / revalidate-on-failure",
            ),
            (
                RegisteredEvidenceSource(
                    source_relative,
                    source_sha256,
                    "supporting-material",
                ),
                RegisteredEvidenceSource(
                    receipt_relative,
                    receipt_sha256,
                    "supporting-material",
                ),
                RegisteredEvidenceSource(
                    catalog_relative,
                    catalog_sha256,
                    "supporting-material",
                ),
            ),
        )

    def _basis_and_graph_with_fixture_evidence(
        self,
        obligations: PreparedObligationSet,
        registrations: tuple[RegisteredEvidenceSource, ...],
    ):  # type: ignore[no-untyped-def]
        manifest = replace(
            self.contract.manifest,
            evidence_sources=(
                *self.contract.manifest.evidence_sources,
                *registrations,
            ),
        )
        receipt = replace(
            self.contract.review_receipt,
            manifest_digest=manifest.digest,
        )
        basis = prepare_reviewer_evidence_basis(
            self.fixture.repo,
            self.compiled,
            manifest,
            receipt,
            obligations,
        )
        graph = replace(
            self.graph,
            source_manifest_digest=manifest.digest,
            obligation_set_digest=obligations.digest,
        )
        return basis, graph

    def _pack(self, *, basis=None, graph=None):  # type: ignore[no-untyped-def]
        basis = basis or self.basis
        graph = graph or self.graph
        plan = build_test_design_plan(graph, context=self.fixture.context)
        cases = plan.deterministic_cases
        draft = render_test_cases(cases, scope_title=self.fixture.context.scope_title)
        gate = validate_suite(
            graph=graph,
            cases=cases,
            markdown=draft,
            checked_path="shadow.md",
        )
        self.assertTrue(gate.passed, gate.to_dict())
        return build_reviewer_evidence_pack(
            basis,
            graph,
            cases,
            draft,
            gate.draft_sha256,
            reviewer_acceptance_contract(schema_version=2),
        )

    def test_exported_basis_reloads_and_requalifies_registered_files(self) -> None:
        document = self.basis.to_document()

        reloaded = load_reviewer_evidence_basis_document(
            self.fixture.repo,
            document,
        )

        self.assertEqual(self.basis.basis_digest, reloaded.basis_digest)
        document["source_manifest"]["assertions"] = document["source_manifest"][
            "assertions"
        ][:-1]
        with self.assertRaises(ReviewerEvidenceError):
            load_reviewer_evidence_basis_document(self.fixture.repo, document)

    def test_literal_rows_cells_and_rows_without_obligations_are_preserved(self) -> None:
        payload = self._pack().to_dict()
        rows = payload["literal_source_evidence"]

        self.assertEqual(self.compiled.baseline.candidate_count, len(rows))
        self.assertTrue(any("только" in item["bounded_source_text"] for item in rows))
        constraint_rows = [
            item for item in rows if item["source_constraint_projection"]
        ]
        self.assertEqual(1, len(constraint_rows))
        self.assertEqual(
            [
                "digits",
                "special-characters-other-than-hyphen",
            ],
            [
                item["negative_class"]
                for item in constraint_rows[0]["source_constraint_projection"]
            ],
        )
        self.assertTrue(
            all(
                item["ui_oracle_derivation"]
                == "forbidden-from-restriction-alone"
                for item in constraint_rows[0]["source_constraint_projection"]
            )
        )
        self.assertTrue(any(item["structured_cells"] for item in rows))
        self.assertTrue(
            all(
                item["section_path_status"]
                == "materialized-from-preceding-xhtml-headings"
                for item in rows
            )
        )
        table_rows = [item for item in rows if item["element_kind"] == "tr"]
        self.assertEqual(2, len(table_rows))
        self.assertEqual(
            {"/*/*[1]/*[2]"},
            {item["table_identity"] for item in table_rows},
        )
        heading_evidence = [
            heading
            for item in rows
            for heading in item["section_heading_evidence"]
        ]
        self.assertTrue(heading_evidence)
        for heading in heading_evidence:
            self.assertTrue(heading["heading_evidence_id"].startswith("XHTML-HEADING-"))
            self.assertEqual(rows[0]["source_path"], heading["source_path"])
            self.assertEqual(
                hashlib.sha256(
                    heading["bounded_source_text"].encode("utf-8")
                ).hexdigest(),
                heading["bounded_source_text_sha256"],
            )
            self.assertEqual(
                rows[0]["source_file_sha256"],
                heading["source_file_sha256"],
            )
        self.assertEqual(
            "materialized-for-all-xhtml-literal-rows",
            payload["source_structure"]["section_path_status"],
        )
        row_ids_with_cases = {
            item["source_row_id"]
            for item in payload["coverage_mapping"]
            if item["case_key"]
        }
        self.assertGreater(len(rows), len(row_ids_with_cases))
        self.assertTrue(
            any(
                not item["obligation_id"]
                for item in payload["coverage_mapping"]
            )
        )
        primary_case_rows = [
            item for item in payload["coverage_mapping"] if item["case_key"]
        ]
        self.assertEqual(len(self.graph.cases), len(primary_case_rows))
        self.assertEqual(
            len(primary_case_rows),
            len({item["case_key"] for item in primary_case_rows}),
        )
        self.assertEqual(
            "complete-literal-registered-artifact",
            payload["coverage_gaps"]["status"],
        )
        self.assertEqual([], payload["coverage_gaps"]["registered_gap_ids"])
        self.assertIn("No gaps.", payload["coverage_gaps"]["content"])

    def test_projection_cannot_drop_any_assertion_including_siblings(self) -> None:
        incomplete = replace(
            self.graph,
            properties=self.graph.properties[:-1],
        )
        plan = build_test_design_plan(self.graph, context=self.fixture.context)
        cases = plan.deterministic_cases
        draft = render_test_cases(cases, scope_title=self.fixture.context.scope_title)

        with self.assertRaisesRegex(
            ReviewerEvidenceError,
            "source-projection-assertion-loss",
        ):
            build_reviewer_evidence_pack(
                self.basis,
                incomplete,
                cases,
                draft,
                hashlib.sha256(draft.encode("utf-8")).hexdigest(),
                reviewer_acceptance_contract(schema_version=2),
            )

    def test_missing_compiled_scope_row_is_rejected_before_pack_creation(self) -> None:
        incomplete_manifest = replace(
            self.contract.manifest,
            source_rows=self.contract.manifest.source_rows[:-1],
        )

        with self.assertRaisesRegex(
            ReviewerEvidenceError,
            "source-row-completeness",
        ):
            prepare_reviewer_evidence_basis(
                self.fixture.repo,
                self.compiled,
                incomplete_manifest,
                self.contract.review_receipt,
                self.fixture.prepared_obligations,
            )

    def test_complete_relevant_dictionary_is_included_and_unbound_one_is_not(self) -> None:
        original = self.fixture.prepared_obligations.obligations[0]
        values = (
            PreparedDictionaryValue(("DICT-001",), "leaf", "Первый"),
            PreparedDictionaryValue(("DICT-001",), "leaf", "Второй"),
        )
        relevant = replace(
            original,
            dictionary_refs=("DICT-001",),
            dictionary_requirements=(
                PreparedDictionaryRequirement(
                    dictionary_id="DICT-001",
                    coverage_mode="all-leaf-values",
                    required_values=values,
                ),
            ),
        )
        unrelated = PreparedObligation(
            obligation_id="OBL-UNRELATED",
            atom_id="ATOM-UNRELATED",
            source_refs=("SRC-001",),
            atomic_statement="Unrelated non-applicable dictionary context.",
            observable_oracle="",
            test_intent="Not applicable.",
            coverage_status="not-applicable",
            gap_id="",
            dictionary_refs=("DICT-UNRELATED",),
            dictionary_requirements=(
                PreparedDictionaryRequirement(
                    dictionary_id="DICT-UNRELATED",
                    coverage_mode="reference-only",
                    fixture_values=(
                        PreparedDictionaryValue(
                            ("DICT-UNRELATED",), "leaf", "Не передавать"
                        ),
                    ),
                ),
            ),
            notes="",
        )
        obligations = PreparedObligationSet.create(
            package_id=self.fixture.prepared_obligations.package_id,
            obligations=(relevant, unrelated),
            coverage_gaps=(),
        )
        basis = prepare_reviewer_evidence_basis(
            self.fixture.repo,
            self.compiled,
            self.contract.manifest,
            self.contract.review_receipt,
            obligations,
        )
        graph = replace(self.graph, obligation_set_digest=obligations.digest)

        dictionaries = self._pack(basis=basis, graph=graph).to_dict()["dictionaries"]

        self.assertEqual(["DICT-001"], [item["dictionary_id"] for item in dictionaries])
        self.assertEqual(
            {"Первый", "Второй"},
            {item["value"] for item in dictionaries[0]["required_values"]},
        )
        self.assertEqual("closed", dictionaries[0]["dictionary_type"])
        self.assertEqual(
            "version-and-effective-date-not-declared-in-qualified-contract",
            dictionaries[0]["metadata_status"],
        )
        self.assertRegex(dictionaries[0]["content_sha256"], r"^[0-9a-f]{64}$")
        self.assertRegex(
            dictionaries[0]["source_binding_sha256"], r"^[0-9a-f]{64}$"
        )
        self.assertEqual(
            [original.obligation_id],
            [
                item["obligation_id"]
                for item in dictionaries[0]["obligation_bindings"]
            ],
        )
        property_id = dictionaries[0]["obligation_bindings"][0]["property_id"]
        expected_property = next(
            item for item in self.graph.properties if item.property_id == property_id
        )
        self.assertEqual(
            expected_property.subject_key,
            dictionaries[0]["obligation_bindings"][0]["subject_key"],
        )

    def test_existing_contact_closed_dictionary_branch_fixtures_are_qualified_subsets(
        self,
    ) -> None:
        project_root = Path(__file__).resolve().parents[1]
        ft_root = project_root / "fts" / "AutoFin"
        registry_path = ft_root / "scope-registry.json"
        if not registry_path.is_file():
            self.skipTest("local AutoFin scope registry is not available")
        registry = load_and_resolve_scope_registry(
            registry_path,
            ft_root,
            scope_id="application-card-contact-persons-repeater",
        )
        compiled = compile_scope_source(
            registry,
            scope_id="application-card-contact-persons-repeater",
            repo_root=project_root,
        )
        prepared_root = (
            ft_root
            / "work"
            / "review-cycles"
            / "contact-persons-action-metadata-v1"
            / "prepared-input"
            / "WP-01"
        )
        obligations = load_obligations(prepared_root / "atomic-obligations.json")
        expected_ids = tuple(
            item.obligation_id
            for item in obligations.obligations
            if item.coverage_status == "testable"
        )
        contract = parse_embedded_source_assertion_contract(
            (prepared_root / "source-evidence.md").read_text(encoding="utf-8"),
            project_root,
            expected_scope_slug="application-card-contact-persons-repeater",
            expected_obligation_ids=expected_ids,
        )
        with patch(
            "test_case_agent.reviewer_evidence._coverage_gap_evidence",
            return_value={"status": "isolated-unrelated-gap-check"},
        ):
            basis = prepare_reviewer_evidence_basis(
                project_root,
                compiled,
                contract.manifest,
                contract.review_receipt,
                obligations,
            )
        relevant_ids = (
            "OBL-CP-004-002",
            "OBL-CP-004-003",
            "OBL-CP-004-005",
        )
        assertion = next(
            item
            for item in contract.manifest.assertions
            if "OBL-CP-004-002" in item.obligation_ids
        )
        property_id = "PROP-CONTACT-DICTIONARY"
        prop = CoverageProperty(
            property_id=property_id,
            assertion_id=assertion.assertion_id,
            property_key="contact.relationship.dictionary",
            subject_key="contact.relationship",
            property_kind="dictionary",
            source_row_id=assertion.source_row_id,
            source_path=assertion.source_path,
            source_locator=assertion.locator,
            source_text_sha256=hashlib.sha256(
                assertion.exact_source_text.encode("utf-8")
            ).hexdigest(),
            canonical_statement=assertion.canonical_statement,
            requirement_codes=assertion.requirement_codes,
            disposition="tc",
        )
        prepared_by_id = {
            item.obligation_id: item for item in obligations.obligations
        }
        graph = CoverageGraph(
            schema_version=1,
            ft_slug="AutoFin",
            scope_slug=contract.manifest.scope_slug,
            source_manifest_digest=contract.manifest.digest,
            obligation_set_digest=obligations.digest,
            properties=(prop,),
            obligations=tuple(
                CoverageObligation(
                    obligation_id=obligation_id,
                    property_id=property_id,
                    atom_id=prepared_by_id[obligation_id].traceability_atom_id,
                    coverage_variant="dictionary",
                    condition_key="always",
                    atomic_statement=prepared_by_id[obligation_id].atomic_statement,
                    observable_oracle=prepared_by_id[obligation_id].observable_oracle,
                    coverage_status="testable",
                    requirement_codes=(),
                    gap_id="",
                    calibration_status="none",
                    validation_trigger="",
                    cleanup_strategy="",
                    source_oracle_id="",
                    fixture_values=(),
                    calibration_question="",
                )
                for obligation_id in relevant_ids
            ),
            cases=(),
            gaps=(),
        )

        item = _dictionary_slice(basis, graph)[0]

        self.assertEqual("DICT-001", item["dictionary_id"])
        self.assertEqual(7, len(item["required_values"]))
        self.assertEqual(
            {"иное", "отец/мать"},
            {value["value"] for value in item["fixture_values"]},
        )
        self.assertEqual(
            "qualified-closed-dictionary-subset",
            item["fixture_provenance"]["evidence_kind"],
        )
        self.assertEqual(
            "qualified-source-membership-verified",
            item["fixture_provenance"]["registration_status"],
        )

    def test_qualified_closed_dictionary_values_may_span_source_rows(self) -> None:
        template = self.contract.manifest.source_rows[0]
        rows = (
            replace(
                template,
                source_row_id="SRC-SPLIT-DICT-001",
                bounded_source_text="Допустимое значение: Альфа.",
            ),
            replace(
                template,
                source_row_id="SRC-SPLIT-DICT-002",
                bounded_source_text="Допустимое значение: Бета.",
            ),
        )
        basis = replace(
            self.basis,
            manifest=replace(self.contract.manifest, source_rows=rows),
        )
        required = (
            PreparedDictionaryValue(("DICT-SPLIT",), "leaf", "Альфа"),
            PreparedDictionaryValue(("DICT-SPLIT",), "leaf", "Бета"),
        )

        evidence = _qualified_closed_fixture_subset_provenance(
            basis,
            dictionary_id="DICT-SPLIT",
            required_values=required,
            fixture_values=(required[0],),
            source_row_ids=tuple(row.source_row_id for row in rows),
        )

        bindings = {
            item["value"]: {
                row["source_row_id"] for row in item["source_rows"]
            }
            for item in evidence["source_value_bindings"]
        }
        self.assertEqual({"SRC-SPLIT-DICT-001"}, bindings["Альфа"])
        self.assertEqual({"SRC-SPLIT-DICT-002"}, bindings["Бета"])

    def test_reference_only_dictionary_is_not_mislabeled_external_dynamic(self) -> None:
        original = self.fixture.prepared_obligations.obligations[0]
        values = (
            PreparedDictionaryValue(
                ("DICT-REFERENCE",),
                "leaf",
                "Актуальное значение",
            ),
        )
        provenance, registrations = self._fixture_provenance(
            "DICT-REFERENCE",
            values,
        )
        requirement = PreparedDictionaryRequirement(
            dictionary_id="DICT-REFERENCE",
            coverage_mode="reference-only",
            fixture_values=values,
            fixture_provenance=provenance,
        )
        relevant = replace(
            original,
            dictionary_refs=("DICT-REFERENCE",),
            dictionary_requirements=(requirement,),
        )
        obligations = PreparedObligationSet.create(
            package_id=self.fixture.prepared_obligations.package_id,
            obligations=(relevant,),
            coverage_gaps=(),
        )
        basis, graph = self._basis_and_graph_with_fixture_evidence(
            obligations,
            registrations,
        )

        item = self._pack(basis=basis, graph=graph).to_dict()["dictionaries"][0]

        self.assertEqual("reference-only-unspecified", item["dictionary_type"])
        self.assertEqual(
            "reference-only-type-not-declared",
            item["dictionary_type_evidence"][0]["contract"],
        )
        self.assertEqual(
            "registered-and-hash-verified",
            item["fixture_provenance"]["registration_status"],
        )

    def test_relevant_fixture_values_without_provenance_fail_closed(self) -> None:
        original = self.fixture.prepared_obligations.obligations[0]
        requirement = PreparedDictionaryRequirement(
            dictionary_id="DICT-UNBOUND",
            coverage_mode="reference-only",
            fixture_values=(
                PreparedDictionaryValue(
                    ("DICT-UNBOUND",),
                    "leaf",
                    "Неподтверждённое значение",
                ),
            ),
        )
        relevant = replace(
            original,
            dictionary_refs=("DICT-UNBOUND",),
            dictionary_requirements=(requirement,),
        )
        obligations = PreparedObligationSet.create(
            package_id=self.fixture.prepared_obligations.package_id,
            obligations=(relevant,),
            coverage_gaps=(),
        )
        basis = prepare_reviewer_evidence_basis(
            self.fixture.repo,
            self.compiled,
            self.contract.manifest,
            self.contract.review_receipt,
            obligations,
        )
        graph = replace(self.graph, obligation_set_digest=obligations.digest)

        with self.assertRaisesRegex(
            ReviewerEvidenceError,
            "dictionary-fixture-provenance-missing",
        ):
            self._pack(basis=basis, graph=graph)

    def test_static_fixture_rejects_cross_dictionary_source_substitution(self) -> None:
        original = self.fixture.prepared_obligations.obligations[0]
        values = (
            PreparedDictionaryValue(("DICT-A",), "leaf", "Value A"),
        )
        provenance, registrations = self._fixture_provenance("DICT-A", values)
        source_path = self.fixture.repo.joinpath(*provenance.source_path.split("/"))
        source_payload = json.loads(source_path.read_text(encoding="utf-8"))
        source_payload["dictionary_id"] = "DICT-WRONG"
        source_path.write_text(
            json.dumps(
                source_payload,
                ensure_ascii=False,
                sort_keys=True,
                separators=(",", ":"),
            ),
            encoding="utf-8",
        )
        source_sha256 = hashlib.sha256(source_path.read_bytes()).hexdigest()
        provenance = replace(provenance, source_sha256=source_sha256)
        registrations = tuple(
            replace(item, sha256=source_sha256)
            if item.path == provenance.source_path
            else item
            for item in registrations
        )
        requirement = PreparedDictionaryRequirement(
            dictionary_id="DICT-A",
            coverage_mode="reference-only",
            fixture_values=values,
            fixture_provenance=provenance,
        )
        relevant = replace(
            original,
            dictionary_refs=("DICT-A",),
            dictionary_requirements=(requirement,),
        )
        obligations = PreparedObligationSet.create(
            package_id=self.fixture.prepared_obligations.package_id,
            obligations=(relevant,),
            coverage_gaps=(),
        )
        basis, graph = self._basis_and_graph_with_fixture_evidence(
            obligations,
            registrations,
        )

        with self.assertRaisesRegex(
            ReviewerEvidenceError,
            "dictionary-fixture-identity-mismatch",
        ):
            self._pack(basis=basis, graph=graph)

    def test_external_dynamic_dictionary_requires_exact_dependency_contract(self) -> None:
        original = self.fixture.prepared_obligations.obligations[0]
        values = (
            PreparedDictionaryValue(
                ("DICT-DYNAMIC",),
                "leaf",
                "Проверенный fixture",
            ),
        )
        provenance, registrations = self._fixture_provenance(
            "DICT-DYNAMIC",
            values,
            external_dynamic=True,
        )
        requirement = PreparedDictionaryRequirement(
            dictionary_id="DICT-DYNAMIC",
            coverage_mode="reference-only",
            fixture_values=values,
            fixture_provenance=provenance,
        )
        relevant = replace(
            original,
            source_refs=(*original.source_refs, "DEP-005"),
            notes="External-dynamic dictionary dependency: DEP-005.",
            dictionary_refs=("DICT-DYNAMIC",),
            dictionary_requirements=(requirement,),
        )
        obligations = PreparedObligationSet.create(
            package_id=self.fixture.prepared_obligations.package_id,
            obligations=(relevant,),
            coverage_gaps=(),
        )
        basis, graph = self._basis_and_graph_with_fixture_evidence(
            obligations,
            registrations,
        )

        item = self._pack(basis=basis, graph=graph).to_dict()["dictionaries"][0]

        self.assertEqual("external-dynamic", item["dictionary_type"])
        self.assertEqual(
            ["DEP-005"],
            item["dictionary_type_evidence"][0]["dependency_ids"],
        )
        self.assertEqual(
            "verified-once",
            item["fixture_provenance"]["verification"]["status"],
        )
        self.assertEqual(
            "prohibited",
            item["fixture_provenance"]["verification"][
                "routine_live_revalidation"
            ],
        )
        self.assertEqual(
            "POST",
            item["fixture_provenance"]["verification"]["request"]["method"],
        )
        self.assertEqual(
            "suggestions-found",
            item["fixture_provenance"]["verification"]["expected_response"][
                "outcome"
            ],
        )

    def test_dadata_receipt_rejects_spoofed_generic_boolean_checks(self) -> None:
        original = self.fixture.prepared_obligations.obligations[0]
        values = (
            PreparedDictionaryValue(
                ("DICT-DADATA-SPOOF",),
                "leaf",
                "Проверенное предложение",
            ),
        )
        provenance, registrations = self._fixture_provenance(
            "DICT-DADATA-SPOOF",
            values,
            external_dynamic=True,
        )
        assert provenance.verification_receipt_path is not None
        receipt_path = self.fixture.repo.joinpath(
            *provenance.verification_receipt_path.split("/")
        )
        receipt_payload = json.loads(receipt_path.read_text(encoding="utf-8"))
        verification = receipt_payload["verification"]
        verification.pop("all_exact_suggestion_matched")
        verification["all_fake"] = True
        verification["attempts"][0].pop("exact_suggestion_matched")
        verification["attempts"][0]["fake"] = True
        receipt_path.write_text(
            json.dumps(
                receipt_payload,
                ensure_ascii=False,
                sort_keys=True,
                separators=(",", ":"),
            ),
            encoding="utf-8",
        )
        receipt_sha256 = hashlib.sha256(receipt_path.read_bytes()).hexdigest()
        provenance = replace(
            provenance,
            verification_receipt_sha256=receipt_sha256,
        )
        registrations = tuple(
            replace(item, sha256=receipt_sha256)
            if item.path == provenance.verification_receipt_path
            else item
            for item in registrations
        )
        requirement = PreparedDictionaryRequirement(
            dictionary_id="DICT-DADATA-SPOOF",
            coverage_mode="reference-only",
            fixture_values=values,
            fixture_provenance=provenance,
        )
        relevant = replace(
            original,
            source_refs=(*original.source_refs, "DEP-DADATA-SPOOF"),
            notes=(
                "External-dynamic dictionary DICT-DADATA-SPOOF dependency: "
                "DEP-DADATA-SPOOF."
            ),
            dictionary_refs=("DICT-DADATA-SPOOF",),
            dictionary_requirements=(requirement,),
        )
        obligations = PreparedObligationSet.create(
            package_id=self.fixture.prepared_obligations.package_id,
            obligations=(relevant,),
            coverage_gaps=(),
        )
        basis, graph = self._basis_and_graph_with_fixture_evidence(
            obligations,
            registrations,
        )

        with self.assertRaisesRegex(
            ReviewerEvidenceError,
            "dictionary-fixture-upstream-receipt-invalid",
        ):
            self._pack(basis=basis, graph=graph)

    def test_existing_autofin_dadata_fixture_is_consumed_without_mutation(self) -> None:
        project_root = Path(__file__).resolve().parents[1]
        actual_root = (
            project_root
            / "fts"
            / "AutoFin"
            / "work"
            / "vendor-references"
        )
        fixture_id = "FX-DADATA-ADDR-POS-001"
        source_name = f"{fixture_id}.response.json"
        receipt_name = f"{fixture_id}.verification.json"
        actual_source = actual_root / "dadata-fixtures" / source_name
        actual_receipt = actual_root / "dadata-fixtures" / receipt_name
        actual_catalog = actual_root / "dadata-fixture-catalog.md"
        self.assertTrue(actual_source.is_file())
        self.assertTrue(actual_receipt.is_file())
        self.assertTrue(actual_catalog.is_file())
        frozen_bytes = {
            actual_source: actual_source.read_bytes(),
            actual_receipt: actual_receipt.read_bytes(),
            actual_catalog: actual_catalog.read_bytes(),
        }

        target_root = self.fixture.ft / "work" / "vendor-references"
        target_fixture_root = target_root / "dadata-fixtures"
        target_fixture_root.mkdir(parents=True, exist_ok=True)
        source = target_fixture_root / source_name
        receipt = target_fixture_root / receipt_name
        catalog = target_root / "dadata-fixture-catalog.md"
        shutil.copyfile(actual_source, source)
        shutil.copyfile(actual_receipt, receipt)
        shutil.copyfile(actual_catalog, catalog)
        receipt_payload = json.loads(receipt.read_text(encoding="utf-8"))
        exact_suggestion = receipt_payload["expected_response"]["exact_suggestion"]
        values = (
            PreparedDictionaryValue(
                ("DICT-DADATA-ADDR",),
                "leaf",
                exact_suggestion,
            ),
        )
        source_relative = self.fixture._relative(source)
        receipt_relative = self.fixture._relative(receipt)
        catalog_relative = self.fixture._relative(catalog)
        source_sha256 = hashlib.sha256(source.read_bytes()).hexdigest()
        receipt_sha256 = hashlib.sha256(receipt.read_bytes()).hexdigest()
        catalog_sha256 = hashlib.sha256(catalog.read_bytes()).hexdigest()
        verified_at = receipt_payload["verification"]["attempts"][-1][
            "checked_at_utc"
        ]
        provenance = PreparedDictionaryFixtureProvenance(
            evidence_kind="external-dynamic-fixture",
            dictionary_id="DICT-DADATA-ADDR",
            fixture_id=fixture_id,
            source_path=source_relative,
            source_locator="json-pointer:",
            source_sha256=source_sha256,
            value_set_sha256=prepared_dictionary_value_set_sha256(values),
            version=str(receipt_payload["schema_version"]),
            effective_date=verified_at[:10],
            provider="DaData",
            verification_receipt_path=receipt_relative,
            verification_receipt_locator="json-pointer:",
            verification_receipt_sha256=receipt_sha256,
            lifecycle_source_path=catalog_relative,
            lifecycle_source_locator=f"markdown-table-row:{fixture_id}",
            lifecycle_source_sha256=catalog_sha256,
            verified_at=verified_at,
            lifecycle_policy="verified-once / revalidate-on-failure",
        )
        registrations = tuple(
            RegisteredEvidenceSource(
                self.fixture._relative(path),
                hashlib.sha256(path.read_bytes()).hexdigest(),
                "supporting-material",
            )
            for path in (source, receipt, catalog)
        )
        requirement = PreparedDictionaryRequirement(
            dictionary_id="DICT-DADATA-ADDR",
            coverage_mode="reference-only",
            fixture_values=values,
            fixture_provenance=provenance,
        )
        original = self.fixture.prepared_obligations.obligations[0]
        relevant = replace(
            original,
            source_refs=(*original.source_refs, "DEP-DADATA-ADDR"),
            notes=(
                "External-dynamic dictionary DICT-DADATA-ADDR dependency: "
                "DEP-DADATA-ADDR."
            ),
            dictionary_refs=("DICT-DADATA-ADDR",),
            dictionary_requirements=(requirement,),
        )
        obligations = PreparedObligationSet.create(
            package_id=self.fixture.prepared_obligations.package_id,
            obligations=(relevant,),
            coverage_gaps=(),
        )
        basis, graph = self._basis_and_graph_with_fixture_evidence(
            obligations,
            registrations,
        )

        item = self._pack(basis=basis, graph=graph).to_dict()["dictionaries"][0]

        verification = item["fixture_provenance"]["verification"]
        self.assertEqual(
            "самара авроры 7 12",
            verification["request"]["parameters"]["query"],
        )
        self.assertEqual(exact_suggestion, verification["expected_response"]["exact_suggestion"])
        self.assertEqual("verified-once", verification["status"])
        self.assertEqual(
            frozen_bytes,
            {
                actual_source: actual_source.read_bytes(),
                actual_receipt: actual_receipt.read_bytes(),
                actual_catalog: actual_catalog.read_bytes(),
            },
        )

    def test_fixture_source_must_be_registered_and_cannot_drift(self) -> None:
        original = self.fixture.prepared_obligations.obligations[0]
        values = (
            PreparedDictionaryValue(
                ("DICT-REGISTERED",),
                "leaf",
                "Замороженное значение",
            ),
        )
        provenance, registrations = self._fixture_provenance(
            "DICT-REGISTERED",
            values,
        )
        requirement = PreparedDictionaryRequirement(
            dictionary_id="DICT-REGISTERED",
            coverage_mode="reference-only",
            fixture_values=values,
            fixture_provenance=provenance,
        )
        relevant = replace(
            original,
            dictionary_refs=("DICT-REGISTERED",),
            dictionary_requirements=(requirement,),
        )
        obligations = PreparedObligationSet.create(
            package_id=self.fixture.prepared_obligations.package_id,
            obligations=(relevant,),
            coverage_gaps=(),
        )
        unregistered_basis = prepare_reviewer_evidence_basis(
            self.fixture.repo,
            self.compiled,
            self.contract.manifest,
            self.contract.review_receipt,
            obligations,
        )
        unregistered_graph = replace(
            self.graph,
            obligation_set_digest=obligations.digest,
        )
        with self.assertRaisesRegex(
            ReviewerEvidenceError,
            "dictionary-fixture-evidence-unregistered",
        ):
            self._pack(basis=unregistered_basis, graph=unregistered_graph)

        basis, graph = self._basis_and_graph_with_fixture_evidence(
            obligations,
            registrations,
        )
        source = self.fixture.repo.joinpath(*provenance.source_path.split("/"))
        source.write_text('{"values":["drift"]}', encoding="utf-8")
        with self.assertRaisesRegex(
            ReviewerEvidenceError,
            "registered-file-drift",
        ):
            self._pack(basis=basis, graph=graph)

    def test_dadata_receipt_must_attest_verified_once_lifecycle(self) -> None:
        original = self.fixture.prepared_obligations.obligations[0]
        values = (
            PreparedDictionaryValue(
                ("DICT-DADATA",),
                "leaf",
                "г Самара, ул Авроры, д 7",
            ),
        )
        provenance, registrations = self._fixture_provenance(
            "DICT-DADATA",
            values,
            external_dynamic=True,
        )
        assert provenance.lifecycle_source_path is not None
        catalog_path = self.fixture.repo.joinpath(
            *provenance.lifecycle_source_path.split("/")
        )
        catalog_path.write_text(
            catalog_path.read_text(encoding="utf-8").replace(
                "verified-once / revalidate-on-failure",
                "always-live-revalidate",
            ),
            encoding="utf-8",
        )
        catalog_sha256 = hashlib.sha256(catalog_path.read_bytes()).hexdigest()
        provenance = replace(
            provenance,
            lifecycle_source_sha256=catalog_sha256,
        )
        registrations = tuple(
            replace(item, sha256=catalog_sha256)
            if item.path == provenance.lifecycle_source_path
            else item
            for item in registrations
        )
        requirement = PreparedDictionaryRequirement(
            dictionary_id="DICT-DADATA",
            coverage_mode="reference-only",
            fixture_values=values,
            fixture_provenance=provenance,
        )
        relevant = replace(
            original,
            source_refs=(*original.source_refs, "DEP-DADATA"),
            notes=(
                "External-dynamic dictionary DICT-DADATA dependency: DEP-DADATA."
            ),
            dictionary_refs=("DICT-DADATA",),
            dictionary_requirements=(requirement,),
        )
        obligations = PreparedObligationSet.create(
            package_id=self.fixture.prepared_obligations.package_id,
            obligations=(relevant,),
            coverage_gaps=(),
        )
        basis, graph = self._basis_and_graph_with_fixture_evidence(
            obligations,
            registrations,
        )

        with self.assertRaisesRegex(
            ReviewerEvidenceError,
            "dictionary-fixture-lifecycle-invalid",
        ):
            self._pack(basis=basis, graph=graph)

    def test_generic_dependency_does_not_classify_multiple_dictionaries(self) -> None:
        original = self.fixture.prepared_obligations.obligations[0]
        registrations: list[RegisteredEvidenceSource] = []

        def requirement(dictionary_id: str) -> PreparedDictionaryRequirement:
            values = (
                PreparedDictionaryValue(
                    (dictionary_id,),
                    "leaf",
                    f"Fixture {dictionary_id}",
                ),
            )
            provenance, current_registrations = self._fixture_provenance(
                dictionary_id,
                values,
            )
            registrations.extend(current_registrations)
            return PreparedDictionaryRequirement(
                dictionary_id=dictionary_id,
                coverage_mode="reference-only",
                fixture_values=values,
                fixture_provenance=provenance,
            )

        relevant = replace(
            original,
            source_refs=(*original.source_refs, "DEP-005"),
            notes="External-dynamic dictionary dependency: DEP-005.",
            dictionary_refs=("DICT-FIRST", "DICT-SECOND"),
            dictionary_requirements=(
                requirement("DICT-FIRST"),
                requirement("DICT-SECOND"),
            ),
        )
        obligations = PreparedObligationSet.create(
            package_id=self.fixture.prepared_obligations.package_id,
            obligations=(relevant,),
            coverage_gaps=(),
        )
        basis, graph = self._basis_and_graph_with_fixture_evidence(
            obligations,
            tuple(registrations),
        )

        dictionaries = self._pack(basis=basis, graph=graph).to_dict()["dictionaries"]

        self.assertEqual(
            {"reference-only-unspecified"},
            {item["dictionary_type"] for item in dictionaries},
        )
        self.assertEqual(
            {"generic-ambiguous"},
            {
                item["dictionary_type_evidence"][0]["dependency_binding"]
                for item in dictionaries
            },
        )

    def test_dictionary_specific_dependency_disambiguates_multiple_dictionaries(self) -> None:
        original = self.fixture.prepared_obligations.obligations[0]
        registrations: list[RegisteredEvidenceSource] = []

        def requirement(dictionary_id: str) -> PreparedDictionaryRequirement:
            values = (
                PreparedDictionaryValue(
                    (dictionary_id,),
                    "leaf",
                    f"Fixture {dictionary_id}",
                ),
            )
            provenance, current_registrations = self._fixture_provenance(
                dictionary_id,
                values,
                external_dynamic=True,
            )
            registrations.extend(current_registrations)
            return PreparedDictionaryRequirement(
                dictionary_id=dictionary_id,
                coverage_mode="reference-only",
                fixture_values=values,
                fixture_provenance=provenance,
            )

        relevant = replace(
            original,
            source_refs=(*original.source_refs, "DEP-005", "DEP-006"),
            notes=(
                "External-dynamic dictionary DICT-FIRST dependency: DEP-005. "
                "External-dynamic dictionary DICT-SECOND dependency: DEP-006."
            ),
            dictionary_refs=("DICT-FIRST", "DICT-SECOND"),
            dictionary_requirements=(
                requirement("DICT-FIRST"),
                requirement("DICT-SECOND"),
            ),
        )
        obligations = PreparedObligationSet.create(
            package_id=self.fixture.prepared_obligations.package_id,
            obligations=(relevant,),
            coverage_gaps=(),
        )
        basis, graph = self._basis_and_graph_with_fixture_evidence(
            obligations,
            tuple(registrations),
        )

        dictionaries = self._pack(basis=basis, graph=graph).to_dict()["dictionaries"]

        self.assertEqual(
            {"external-dynamic"},
            {item["dictionary_type"] for item in dictionaries},
        )
        by_id = {item["dictionary_id"]: item for item in dictionaries}
        self.assertEqual(
            ["DEP-005"],
            by_id["DICT-FIRST"]["dictionary_type_evidence"][0]["dependency_ids"],
        )
        self.assertEqual(
            ["DEP-006"],
            by_id["DICT-SECOND"]["dictionary_type_evidence"][0]["dependency_ids"],
        )

    def test_supporting_cross_row_binding_is_materialized_to_the_full_chain(self) -> None:
        primary = self.contract.manifest.assertions[1]
        support_row = self.contract.manifest.source_rows[0]
        assertion = replace(
            primary,
            supporting_source_bindings=(
                SupportingSourceBinding(
                    source_row_id=support_row.source_row_id,
                    evidence_role="constraint",
                    exact_source_fragment=support_row.bounded_source_text,
                ),
            ),
        )
        assertions = (
            self.contract.manifest.assertions[0],
            assertion,
            self.contract.manifest.assertions[2],
        )
        manifest = replace(self.contract.manifest, assertions=assertions)
        receipt = replace(
            self.contract.review_receipt,
            manifest_digest=manifest.digest,
        )
        basis = prepare_reviewer_evidence_basis(
            self.fixture.repo,
            self.compiled,
            manifest,
            receipt,
            self.fixture.prepared_obligations,
        )
        graph = replace(self.graph, source_manifest_digest=manifest.digest)

        mapping = self._pack(basis=basis, graph=graph).to_dict()[
            "supporting_evidence_mapping"
        ]

        self.assertEqual(1, len(mapping))
        self.assertEqual(support_row.source_row_id, mapping[0]["source_row_id"])
        self.assertEqual(primary.source_row_id, mapping[0]["primary_source_row_id"])
        self.assertEqual(primary.assertion_id, mapping[0]["assertion_id"])
        self.assertTrue(mapping[0]["case_key"])
        self.assertTrue(mapping[0]["tc_id"])

    def test_repeater_sibling_actions_have_role_tagged_design_support_chains(self) -> None:
        from tests.test_test_design import _repeater_graph_and_context

        graph, context = _repeater_graph_and_context()
        cases = build_test_design_plan(graph, context=context).deterministic_cases

        mapping = build_design_support_mapping(graph, cases)

        add_case = "customer|customer-name|source-add-row|repeater-add|always"
        delete_case = "customer|delete-control|source-delete-row|repeater-delete|always"
        visibility_case = "customer|customer-name|visibility|positive|always"
        self.assertEqual(
            [("cleanup", "OBL-DELETE")],
            [
                (item["support_role"], item["obligation_id"])
                for item in mapping
                if item["case_key"] == add_case
            ],
        )
        self.assertEqual(
            [("setup", "OBL-ADD"), ("setup", "OBL-ADD")],
            [
                (item["support_role"], item["obligation_id"])
                for item in mapping
                if item["case_key"] == delete_case
            ],
        )
        self.assertEqual(
            {("action", "OBL-ADD"), ("cleanup", "OBL-DELETE")},
            {
                (item["support_role"], item["obligation_id"])
                for item in mapping
                if item["case_key"] == visibility_case
            },
        )
        self.assertTrue(
            all(item["obligation_id"] in next(
                case.traceability for case in cases if case.case_key == item["case_key"]
            ) for item in mapping)
        )

    def test_scope_gap_ids_must_be_materialized_in_registered_artifact(self) -> None:
        compiled = replace(
            self.compiled,
            definition=replace(
                self.compiled.definition,
                gap_ids=("GAP-MISSING",),
            ),
        )

        with self.assertRaisesRegex(
            ReviewerEvidenceError,
            "coverage-gap-id-mismatch",
        ):
            prepare_reviewer_evidence_basis(
                self.fixture.repo,
                compiled,
                self.contract.manifest,
                self.contract.review_receipt,
                self.fixture.prepared_obligations,
            )

    def test_missing_scope_mockup_registration_is_rejected(self) -> None:
        compiled = replace(
            self.compiled,
            definition=replace(
                self.compiled.definition,
                reference_paths=("mockups/screen.png",),
            ),
        )

        with self.assertRaisesRegex(
            ReviewerEvidenceError,
            "scope-mockup-set-mismatch",
        ):
            prepare_reviewer_evidence_basis(
                self.fixture.repo,
                compiled,
                self.contract.manifest,
                self.contract.review_receipt,
                self.fixture.prepared_obligations,
            )

    def test_registered_mockup_drift_after_basis_creation_is_rejected(self) -> None:
        mockup = self.fixture.ft / "mockups" / "screen.png"
        mockup.parent.mkdir()
        Image.new("RGB", (3, 2), color=(31, 97, 173)).save(mockup)
        mockup_relative = self.fixture._relative(mockup)
        manifest = build_source_assertion_manifest(
            self.fixture.repo,
            scope_slug=self.contract.manifest.scope_slug,
            coverage_gaps_path=self.contract.manifest.coverage_gaps_artifact.path,
            source_paths=tuple(item.path for item in self.contract.manifest.sources),
            assertions=self.contract.manifest.assertions,
            source_row_extraction_spec_digest=(
                self.contract.manifest.source_row_extraction_spec_digest
            ),
            source_row_baseline_digest=(
                self.contract.manifest.source_row_baseline_digest
            ),
            source_row_candidate_count=(
                self.contract.manifest.source_row_candidate_count
            ),
            source_rows=self.contract.manifest.source_rows,
            evidence_sources=tuple(
                (item.path, item.role)
                for item in self.contract.manifest.evidence_sources
            ),
            clarifications=self.contract.manifest.clarifications,
            mockups=((mockup_relative, "Main screen", ("state:maximum",)),),
            expected_source_rows=self.contract.manifest.source_rows,
        )
        receipt = replace(
            self.contract.review_receipt,
            manifest_digest=manifest.digest,
        )
        receipt.validate(manifest)
        compiled = replace(
            self.compiled,
            definition=replace(
                self.compiled.definition,
                reference_paths=("mockups/screen.png",),
            ),
        )
        basis = prepare_reviewer_evidence_basis(
            self.fixture.repo,
            compiled,
            manifest,
            receipt,
            self.fixture.prepared_obligations,
        )
        Image.new("RGB", (3, 2), color=(173, 31, 97)).save(mockup)

        with self.assertRaisesRegex(ReviewerEvidenceError, "registered-file-drift"):
            basis.verify()

    def test_hash_current_docx_xhtml_semantic_mismatch_blocks_pack(self) -> None:
        document = Document(self.fixture.docx)
        document.tables[0].cell(1, 0).text = "Semantically different field rule."
        document.save(self.fixture.docx)
        manifest = build_source_assertion_manifest(
            self.fixture.repo,
            scope_slug=self.contract.manifest.scope_slug,
            coverage_gaps_path=self.contract.manifest.coverage_gaps_artifact.path,
            source_paths=tuple(item.path for item in self.contract.manifest.sources),
            assertions=self.contract.manifest.assertions,
            source_row_extraction_spec_digest=(
                self.contract.manifest.source_row_extraction_spec_digest
            ),
            source_row_baseline_digest=(
                self.contract.manifest.source_row_baseline_digest
            ),
            source_row_candidate_count=(
                self.contract.manifest.source_row_candidate_count
            ),
            source_rows=self.contract.manifest.source_rows,
            evidence_sources=tuple(
                (item.path, item.role)
                for item in self.contract.manifest.evidence_sources
            ),
            clarifications=self.contract.manifest.clarifications,
            mockups=tuple(
                (
                    item.path,
                    item.screen_name,
                    item.locators,
                )
                for item in self.contract.manifest.mockups
            ),
            expected_source_rows=self.contract.manifest.source_rows,
        )
        receipt = replace(
            self.contract.review_receipt,
            manifest_digest=manifest.digest,
        )
        basis = prepare_reviewer_evidence_basis(
            self.fixture.repo,
            self.compiled,
            manifest,
            receipt,
            self.fixture.prepared_obligations,
        )
        graph = replace(self.graph, source_manifest_digest=manifest.digest)

        with self.assertRaisesRegex(
            ReviewerEvidenceError,
            "source-parity-docx-xhtml-table-row-mismatch",
        ):
            self._pack(basis=basis, graph=graph)


if __name__ == "__main__":
    unittest.main()
